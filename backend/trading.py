"""Catalog-backed checkout and derived receivable schedule, under tenant RLS."""
from calendar import monthrange
from datetime import date
from decimal import Decimal
from uuid import UUID
from pydantic import Field, model_validator
from fastapi import Depends, HTTPException
from .workspace import router, Input, Session, session, Sale, MONEY_MAX
from .finance import simulate

class Line(Input):
    product_id: UUID
    quantity: int = Field(ge=1,le=1000000)
    unit_price: int | None = Field(default=None,gt=0,le=MONEY_MAX)

class Checkout(Input):
    reference: UUID
    sold_on: date
    first_due_on: date
    lines: list[Line] = Field(min_length=1,max_length=50)
    terms_id: UUID | None = None
    installments: int = Field(default=1,ge=1,le=12)
    manage_stock: bool = True
    received: bool = False
    tax_rate: Decimal = Field(default=0,ge=0,le=50)
    commission_rate: Decimal = Field(default=0,ge=0,le=50)
    card_base: Decimal = Field(default=0,ge=0,le=50)
    anticipation_rate: Decimal = Field(default=0,ge=0,le=20)

    @model_validator(mode='after')
    def dates(self):
        if self.first_due_on<self.sold_on: raise ValueError('Vencimento anterior à venda.')
        return self

def quote(db,cid,body):
    db.company(cid)
    rates={k:getattr(body,k) for k in ('tax_rate','commission_rate','card_base','anticipation_rate')}
    installments=body.installments
    if body.terms_id:
        terms=db.rows('payment_terms',cid,id='eq.'+str(body.terms_id))
        if not terms: raise HTTPException(404,'Condição não encontrada nesta empresa.')
        rates={k:Decimal(str(terms[0][k])) for k in rates};installments=terms[0]['installments']
    products={p['id']:p for p in db.rows('products',cid)}
    result=[]
    for i,line in enumerate(body.lines):
        p=products.get(str(line.product_id))
        if not p or not p['active']: raise HTTPException(422,'Selecione produtos ativos da sua empresa.')
        price=p['unit_price'] if line.unit_price is None else line.unit_price
        totals=simulate(Decimal(price*line.quantity)/100,Decimal(p['unit_cost']*line.quantity)/100,installments,**rates)
        if totals['card']>totals['price']: raise HTTPException(422,'A taxa do cartão supera o valor da venda.')
        record=Sale(external_id=f'{body.reference}:{i+1}',sold_on=body.sold_on,product=p['name'],category=p['category'],quantity=line.quantity,revenue=totals['price'],cmv=totals['cost'],tax=totals['tax'],card=totals['card'],commission=totals['commission'],installments=installments).model_dump(mode='json')
        result.append({**record,'product_id':p['id'],'batch_reference':str(body.reference),'first_due_on':str(body.first_due_on),'stock_managed':body.manage_stock})
    return result

@router.post('/{company_id}/sales/quote')
def sale_quote(company_id:UUID,body:Checkout,db:Session=Depends(session)):
    rows=quote(db,str(company_id),body)
    totals={k:sum(r[k] for r in rows) for k in ('revenue','cmv','tax','card','commission')}
    totals['net']=totals['revenue']-sum(totals[k] for k in ('cmv','tax','card','commission'))
    return {'rows':rows,'totals':totals}

@router.post('/{company_id}/sales/checkout',status_code=201)
def checkout(company_id:UUID,body:Checkout,db:Session=Depends(session)):
    cid=str(company_id)
    with db.transaction() as tx:
        rows=quote(tx,cid,body)
        for pid in sorted({r['product_id'] for r in rows}):
            tx.connection.execute('select pg_advisory_xact_lock(hashtextextended(%s,1))',(pid,))
        if body.manage_stock:
            movements=tx.rows('stock_movements',cid)
            for pid in {r['product_id'] for r in rows}:
                available=sum(m['quantity']*(1 if m['direction']=='in' else -1) for m in movements if m['product_id']==pid)
                required=sum(r['quantity'] for r in rows if r['product_id']==pid)
                if available<required:
                    product=next(r['product'] for r in rows if r['product_id']==pid)
                    raise HTTPException(422,f'Estoque insuficiente para {product}: saldo {available}, venda {required}. Registre a entrada real antes de vender.')
        result=[]
        for r in rows:
            saved=tx.insert('sales',{'company_id':cid,**r});result.append(saved)
            if body.manage_stock:
                tx.insert('stock_movements',{'company_id':cid,'product_id':r['product_id'],'occurred_on':r['sold_on'],'direction':'out','quantity':r['quantity'],'reference':'sale:'+saved['id'],'description':'Venda '+r['external_id']})
            expected=r['revenue']-r['card']
            if body.received and expected>0:
                tx.insert('settlements',{'company_id':cid,'sale_id':saved['id'],'reference':'sale:'+saved['id'],'received_on':r['sold_on'],'amount':expected})
    return {'sales':result}

class Return(Input):
    reference: UUID
    quantity: int = Field(gt=0,le=1000000)
    occurred_on: date
    restock: bool = True
    description: str = Field(min_length=2,max_length=200)

@router.post('/{company_id}/sales/{sale_id}/return',status_code=201)
def return_sale(company_id:UUID,sale_id:UUID,body:Return,db:Session=Depends(session)):
    cid=str(company_id);sid=str(sale_id)
    with db.transaction() as tx:
        tx.company(cid)
        tx.connection.execute('select pg_advisory_xact_lock(hashtextextended(%s,0))',(sid,))
        rows=tx.rows('sales',cid,id='eq.'+sid)
        if not rows or not rows[0].get('product_id'): raise HTTPException(404,'Venda integrada não encontrada.')
        sale=rows[0]
        if str(body.occurred_on)<sale['sold_on']: raise HTTPException(422,'A devolução não pode ser anterior à venda.')
        previous=tx.rows('adjustments',cid,sale_id='eq.'+sid)
        returned=sum(a['returned_quantity'] for a in previous)
        if returned+body.quantity>sale['quantity']: raise HTTPException(422,'Quantidade maior que o saldo ainda não devolvido.')
        amount=sale['revenue']*body.quantity//sale['quantity']
        cost=sale['cmv']*body.quantity//sale['quantity'] if body.restock else 0
        saved=tx.insert('adjustments',{'company_id':cid,'sale_id':sid,'occurred_on':str(body.occurred_on),'kind':'refund','amount':amount,'cost_recovered':cost,'description':body.description,'returned_quantity':body.quantity,'reference':str(body.reference)})
        if body.restock and sale['stock_managed']:
            tx.insert('stock_movements',{'company_id':cid,'product_id':sale['product_id'],'occurred_on':str(body.occurred_on),'direction':'in','quantity':body.quantity,'reference':'return:'+saved['id'],'description':body.description})
    return saved

def add_months(value,n):
    total=value.year*12+value.month-1+n;y,m=divmod(total,12)
    return date(y,m+1,min(value.day,monthrange(y,m+1)[1]))

@router.get('/{company_id}/receivables')
def receivables(company_id:UUID,db:Session=Depends(session)):
    cid=str(company_id);db.company(cid)
    payments=db.rows('settlements',cid);adjustments=db.rows('adjustments',cid);result=[]
    for sale in db.rows('sales',cid):
        refunds=sum(a['amount'] for a in adjustments if a['sale_id']==sale['id'])
        expected=max(0,sale['revenue']-sale['card']-refunds)
        received=sum(p['amount'] for p in payments if p['sale_id']==sale['id'])
        remaining=received;parts=[]
        for i in range(sale['installments']):
            amount=expected//sale['installments']+(1 if i<expected%sale['installments'] else 0)
            paid=min(amount,remaining);remaining-=paid
            due=str(add_months(date.fromisoformat(sale['first_due_on']),i)) if sale.get('first_due_on') else None
            parts.append({'number':i+1,'due_on':due,'amount':amount,'paid':paid,'balance':amount-paid})
        result.append({**sale,'expected':expected,'received':received,'balance':max(0,expected-received),'excess':max(0,received-expected),'returned_quantity':sum(a.get('returned_quantity',0) for a in adjustments if a['sale_id']==sale['id']),'parts':parts})
    return result

class StockLimits(Input):
    min_stock: int = Field(ge=0,le=1000000)
    max_stock: int = Field(ge=0,le=1000000)
    @model_validator(mode='after')
    def bounds(self):
        if self.max_stock<self.min_stock: raise ValueError('Máximo deve ser maior ou igual ao mínimo.')
        return self

@router.patch('/{company_id}/inventory/{product_id}/limits')
def limits(company_id:UUID,product_id:UUID,body:StockLimits,db:Session=Depends(session)):
    cid=str(company_id);db.company(cid)
    rows=db.request('PATCH','products',params={'company_id':'eq.'+cid,'id':'eq.'+str(product_id)},payload=body.model_dump())
    if not rows: raise HTTPException(404,'Produto não encontrado.')
    return rows[0]

@router.get('/{company_id}/inventory/export')
def export_inventory(company_id:UUID,format:str='xlsx',db:Session=Depends(session)):
    from io import BytesIO
    from fastapi import Response
    from .workspace import inventory
    if format not in ('xlsx','pdf'): raise HTTPException(422,'Escolha XLSX ou PDF.')
    with db.transaction(snapshot=True) as tx:
        company=tx.company(str(company_id))
        rows=inventory(company_id,tx)
    headers=['Produto','SKU','Categoria','Quantidade (un.)','Custo unitário (R$)','Valor ao custo (R$)','Mínimo','Máximo']
    values=[[r['name'],r['sku'],r['category'],r['balance'],r['unit_cost']/100,r['balance']*r['unit_cost']/100,r['min_stock'],r['max_stock']] for r in rows]
    output=BytesIO()
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        from openpyxl.utils import get_column_letter
        book=Workbook();sheet=book.active;sheet.title='Estoque atual';sheet.append(headers)
        for row in values:
            sheet.append(row)
            for c in sheet[sheet.max_row][:3]: c.data_type='s'
        for c in sheet[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='28583E')
        for i,width in enumerate([36,20,24,20,24,24,14,14],1):sheet.column_dimensions[get_column_letter(i)].width=width
        for row in sheet.iter_rows(min_row=2):
            row[4].number_format=row[5].number_format='"R$" #,##0.00'
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions;book.save(output)
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from xml.sax.saxutils import escape
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib import colors
        styles=getSampleStyleSheet();styles['BodyText'].fontSize=8
        table_rows=[[Paragraph(escape(str(v)),styles['BodyText']) for v in row] for row in [headers,*values]]
        table=Table(table_rows,colWidths=[145,65,100,75,90,90,65,65],repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf3e5')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,-1),.3,colors.lightgrey)]))
        SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=22,rightMargin=22).build([Paragraph('Estoque — '+escape(company['name']),styles['Title']),Paragraph('Posição atual. Quantidades em unidades; valorização pelo custo atual do cadastro.',styles['BodyText']),Spacer(1,15),table])
        mime='application/pdf'
    return Response(output.getvalue(),media_type=mime,headers={'Content-Disposition':f'attachment; filename="estoque-lucra.{format}"'})
