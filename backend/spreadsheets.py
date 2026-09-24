"""Bounded, preview-first spreadsheet imports; all writes use the tenant RLS session."""
import hashlib
import base64
import binascii
import csv
import io
import zipfile
import json
import re
import unicodedata
from typing import Literal
from pydantic import ValidationError
from datetime import date, datetime
from decimal import Decimal
from uuid import UUID

from fastapi import Depends, HTTPException
from pydantic import Field
from .workspace import router, Input, Product, Sale, Expense, Session, session, MONEY_MAX

class ExpenseImport(Expense):
    external_id: str = Field(min_length=1,max_length=120)

MODELS = {'products':Product,'sales':Sale,'expenses':ExpenseImport}
MONEY = {'unit_cost','unit_price','revenue','cmv','tax','card','commission','amount'}

def normalized(value):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD',str(value)).encode('ascii','ignore').decode().lower())

ALIASES = {
 'sku':['código','cód','referência','código produto','sku'],
 'name':['produto','nome','descrição','nome do produto'],
 'category':['categoria','grupo','departamento'],
 'unit_cost':['custo','custo aquisição','preço de custo','custo unitário'],
 'unit_price':['preço','preço venda','preço de venda','valor unitário'],
 'notes':['observações','obs'], 'active':['ativo'],
 'external_id':['id','referência','identificador','número da venda','código venda','id venda'],
 'sold_on':['data','data venda','data da venda'],
 'product':['produto','nome do produto','produto auto'],
 'quantity':['quantidade','qtd','qtde'],
 'revenue':['receita','receita total','total da venda','valor total'],
 'cmv':['cmv','custo mercadoria','custo mercadoria total'],
 'tax':['imposto','impostos','valor imposto'],
 'card':['taxa cartão','valor taxa cartão'],
 'commission':['comissão','valor comissão'],
 'installments':['parcelas','número de parcelas'],
 'description':['descrição','histórico','despesa'],
 'incurred_on':['data','competência','data competência','data vencimento'],
 'amount':['valor','total','valor despesa'],
}

def suggest_mapping(kind,headers):
    result={}
    for field in MODELS[kind].model_fields:
        matches=[h for h in headers if not (field in MONEY and '%' in h) and normalized(h) in {normalized(v) for v in [field,*ALIASES.get(field,[])]}]
        if len(matches)==1: result[field]=matches[0]
    return result

class ImportOptions(Input):
    header_row: int | None = Field(default=None,ge=1,le=50)
    number_format: Literal['auto','br','us'] = 'auto'
    date_format: Literal['dmy','mdy'] = 'dmy'
    encoding: Literal['utf-8-sig','cp1252'] = 'utf-8-sig'
    delimiter: Literal['auto',',',';','\t'] = 'auto'
    defaults: dict[str,str] = Field(default_factory=dict,max_length=40)
    generate_ids: bool = False

class Spreadsheet(ImportOptions):
    kind: str
    source_id: UUID | None = None
    filename: str = Field(max_length=180)
    content: str = Field(max_length=2800000)
    sheet: str | None = None
    mapping: dict[str,str] = Field(default_factory=dict,max_length=40)
    corrections: dict[int,dict[str,str]] = Field(default_factory=dict,max_length=1000)
    excluded_rows: list[int] = Field(default_factory=list,max_length=1000)

class SourceRow(dict):
    def __init__(self,values,line):
        super().__init__(values)
        self.line=line

def money_value(value,locale):
    text=str(value).strip().replace('R$','').replace('$','').replace(' ','').replace('\xa0','')
    if isinstance(value,str):
        if locale=='auto':
            if re.fullmatch(r'[+-]?\d+[.,]\d{3}',text):
                raise ValueError('Número ambíguo: escolha o formato brasileiro ou internacional.')
            locale='br' if ',' in text and ('.' not in text or text.rfind(',')>text.rfind('.')) else 'us'
        pattern=r'\d+(?:,\d{1,2})?|\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?' if locale=='br' else r'\d+(?:\.\d{1,2})?|\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?'
        if not re.fullmatch(pattern,text): raise ValueError('Valor monetário inválido para o formato escolhido.')
        text=text.replace('.','').replace(',','.') if locale=='br' else text.replace(',','')
    amount=Decimal(text)
    if not amount.is_finite() or amount<0 or amount>MONEY_MAX/100 or amount.as_tuple().exponent < -2:
        raise ValueError('Valor monetário inválido.')
    return int(amount*100)


def read_sheet(body, discover=False):
    if body.kind not in MODELS:
        raise HTTPException(422,'Escolha produtos, vendas ou despesas.')
    try:
        raw=base64.b64decode(body.content,validate=True)
        if len(raw)>2000000: raise ValueError('Arquivo maior que 2 MB.')
        sheets=[]
        if body.filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(f.file_size for f in archive.infolist())>20000000 or len(archive.infolist())>200:
                    raise ValueError('Planilha muito grande quando descompactada.')
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            cached=None
            try:
                sheets=book.sheetnames
                if discover: return [],[],sheets
                sheet=book[body.sheet or sheets[0]]
                sheet.reset_dimensions()
                cached=load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
                cached_sheet=cached[sheet.title]
                cached_sheet.reset_dimensions()
                data=[]
                for index,(row,values_row) in enumerate(zip(sheet.iter_rows(max_col=41),cached_sheet.iter_rows(max_col=41))):
                    if index>1050: raise ValueError("Limite de 1.000 linhas após o cabeçalho.")
                    values=[(v.value if v.value is not None and v.data_type!='e' else f'Fórmula sem resultado: {c.coordinate}') if c.data_type=='f' else c.value for c,v in zip(row,values_row)]
                    while values and values[-1] is None: values.pop()
                    data.append(values)
            finally:
                book.close()
                if cached: cached.close()
        elif body.filename.lower().endswith(('.csv','.tsv')):
            text=raw.decode(body.encoding)
            try: dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
            except csv.Error:
                candidates={d:list(csv.reader(io.StringIO(text[:8192]),delimiter=d))[:50] for d in [',',';','\t']}
                chosen=max(candidates,key=lambda d:max((len(suggest_mapping(body.kind,r)) for r in candidates[d]),default=0))
                dialect=type('DetectedDialect',(csv.excel,),{'delimiter':chosen})
            data=[]
            reader=csv.reader(io.StringIO(text),dialect) if body.delimiter=='auto' else csv.reader(io.StringIO(text),delimiter=body.delimiter)
            for row in reader:
                data.append(row)
                if len(data)>1051: raise ValueError('Limite de linhas excedido.')
        else: raise ValueError('Use CSV UTF-8 ou Excel .xlsx.')
        if len(data)<2: raise ValueError('Arquivo sem registros.')
        candidates=[i for i,r in enumerate(data[:50]) if any(v is not None and str(v).strip() for v in r)]
        start=body.header_row-1 if body.header_row else max(candidates,key=lambda i:len(suggest_mapping(body.kind,[str(v or '') for v in data[i]])),default=0)
        headers=[str(v or '').strip() for v in data[start]]
        body.header_row=start+1
        numbered=[(i+1,r) for i,r in enumerate(data) if i>start and any(v is not None and str(v).strip() for v in r)]
        if len(numbered)>1000: raise ValueError('Limite de 1.000 registros por arquivo.')
        if len(headers)>40 or not all(headers) or len(set(headers))!=len(headers):
            raise ValueError('Use até 40 colunas, com cabeçalhos únicos e preenchidos.')
        if any(len(r)>len(headers) for _,r in numbered): raise ValueError('Há linhas com mais colunas que o cabeçalho.')
        rows=[SourceRow(dict(zip(headers,r)),line) for line,r in numbered]
        return headers,rows,sheets
    except (ValueError,KeyError,IndexError,UnicodeError,zipfile.BadZipFile,binascii.Error,csv.Error) as e:
        raise HTTPException(422,str(e))
    except Exception:
        raise HTTPException(422,'Não foi possível ler o arquivo. Confira se é um CSV UTF-8 ou XLSX válido.') from None


def prepare_rows(body):
    headers,rows,_=read_sheet(body)
    model=MODELS[body.kind]
    mapping=body.mapping or suggest_mapping(body.kind,headers)
    fields=set(model.model_fields)
    if any(k not in fields or v not in headers for k,v in mapping.items()) or any(k not in fields for k in body.defaults):
        raise HTTPException(422,'Mapeamento ou valores fixos inválidos.')
    if any(k not in fields for edit in body.corrections.values() for k in edit):
        raise HTTPException(422,'Campo de correção inválido.')
    identity='sku' if body.kind=='products' else 'external_id'
    occurrences={};prepared=[]
    for row in rows:
        signature=json.dumps(dict(row),sort_keys=True,default=str,ensure_ascii=False)
        digest=hashlib.sha256((body.kind+'|'+str(body.source_id or '')+'|'+signature).encode()).hexdigest()[:32]
        occurrences[digest]=occurrences.get(digest,0)+1
        if row.line in body.excluded_rows: continue
        values={k:row.get(v) for k,v in mapping.items()}
        for k,v in body.defaults.items():
            if values.get(k) is None or values.get(k)=='': values[k]=v
        if body.generate_ids and not values.get(identity): values[identity]='AUTO-'+digest+'-'+str(occurrences[digest])
        values.update(body.corrections.get(row.line,{}))
        prepared.append((row.line,values))
    return prepared

def parse_sheet(body):
    model=MODELS[body.kind];parsed,errors,seen=[],[],set()
    identity='sku' if body.kind=='products' else 'external_id'
    for number,original in prepare_rows(body):
        values=dict(original);field=''
        try:
            for field,v in list(values.items()):
                if isinstance(v,str) and v.startswith('Fórmula sem resultado:'):
                    raise ValueError(v+'. Recalcule no Excel ou informe o valor na revisão.')
                if model.model_fields[field].annotation is str and isinstance(v,(int,float)):
                    values[field]=str(int(v)) if int(v)==v else str(v)
                if field in MONEY: values[field]=money_value(v,body.number_format)
                if field.endswith('_on'):
                    if isinstance(v,datetime): values[field]=v.date().isoformat()
                    elif isinstance(v,date): values[field]=v.isoformat()
                    elif '/' in str(v): values[field]=datetime.strptime(str(v),'%d/%m/%Y' if body.date_format=='dmy' else '%m/%d/%Y').date().isoformat()
                if field=='active' and isinstance(v,str) and normalized(v) in ('sim','nao','ativo','inativo'):
                    values[field]=normalized(v) in ('sim','ativo')
            record=model.model_validate(values).model_dump(mode='json')
            if record[identity] in seen:
                field=identity;raise ValueError('Identificador repetido neste arquivo.')
            seen.add(record[identity]);parsed.append(record)
        except ValidationError as e:
            errors.extend({'line':number,'field':str(item['loc'][0]),'message':'Campo obrigatório ausente ou valor inválido.'} for item in e.errors())
        except (ValueError,TypeError,ArithmeticError) as e:
            errors.append({'line':number,'field':field,'message':str(e) or 'Valor inválido.'})
    return parsed,errors

@router.post('/{company_id}/spreadsheets/read')
def inspect(company_id:UUID,body:Spreadsheet,db:Session=Depends(session)):
    db.company(str(company_id))
    headers,rows,sheets=read_sheet(body,discover=body.filename.lower().endswith('.xlsx') and not body.sheet)
    return {'headers':headers,'sample':rows[:5],'sheets':sheets,'count':len(rows),
            'header_row':body.header_row,'suggested_mapping':suggest_mapping(body.kind,headers),
            'fields':[{'key':k,'required':v.is_required()} for k,v in MODELS[body.kind].model_fields.items()]}

@router.post('/{company_id}/spreadsheets/{operation}')
def import_sheet(company_id:UUID,operation:str,body:Spreadsheet,db:Session=Depends(session)):
    if operation not in ('preview','confirm'): raise HTTPException(404)
    db.company(str(company_id))
    cid=str(company_id)
    source_id=str(body.source_id) if body.source_id else None
    if source_id and not db.rows('integration_sources',cid,id='eq.'+source_id):
        raise HTTPException(404,'Origem não encontrada nesta empresa.')
    audit={'company_id':cid,'source_id':source_id,'kind':body.kind,
           'filename':body.filename.replace('\\','/').split('/')[-1],
           'fingerprint':hashlib.sha256(body.content.encode()).hexdigest()}
    try:
        parsed,errors=parse_sheet(body)
        identity='sku' if body.kind=='products' else 'external_id'
        existing={r.get(identity) for r in db.rows(body.kind,cid)}
        duplicates=[r[identity] for r in parsed if r[identity] in existing]
        valid=bool(parsed) and not errors and not duplicates
        if operation=='confirm':
            if not valid:
                raise HTTPException(422,'Corrija os erros ou identificadores duplicados antes de importar.')
            with db.transaction() as tx:
                job=tx.insert('import_jobs',{**audit,'status':'imported','imported':len(parsed)})
                saved=tx.request('POST',body.kind,payload=[{**r,'company_id':cid} for r in parsed])
                if source_id:
                    tx.request('POST','source_records',payload=[{'company_id':cid,'source_id':source_id,
                        'job_id':job['id'],'kind':body.kind,'external_id':r[identity],'record_id':r['id']} for r in saved])
            return {'imported':len(parsed),'job_id':job['id']}
        if not valid:
            db.insert('import_jobs',{**audit,'status':'rejected','rejected':len(errors)+len(duplicates),
                                     'detail':'Validação: corrija campos inválidos ou identificadores repetidos.'})
        return {'rows':parsed,'errors':errors,'duplicates':duplicates,'can_import':valid,
                'review':[{'line':line,'values':values} for line,values in prepare_rows(body)]}
    except HTTPException as error:
        if error.status_code in (409,413,422):
            db.insert('import_jobs',{**audit,'status':'rejected','rejected':1,
                                     'detail':str(error.detail)[:500]})
        raise
